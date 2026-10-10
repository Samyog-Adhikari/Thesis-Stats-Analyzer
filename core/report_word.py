"""Word report generation for standardized analysis results."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document
from docx.shared import Inches


def results_to_word(results: list[dict], title: str = "Thesis Statistical Analysis Report") -> bytes:
    """Create an in-memory DOCX report containing result tables and figures."""
    document = Document(); document.add_heading(title, level=0); document.add_paragraph("Generated locally. Uploaded research data is not stored by this app.")
    for result in results:
        document.add_heading(result.get("title", "Analysis"), level=1); document.add_paragraph(f"N = {result.get('n', 'NA')}")
        for warning in result.get("warnings", []): document.add_paragraph(f"Warning: {warning}", style="Intense Quote")
        for table in result.get("tables", []):
            if not isinstance(table, pd.DataFrame): continue
            doc_table = document.add_table(rows=1, cols=len(table.columns) + (1 if table.index.name or not isinstance(table.index, pd.RangeIndex) else 0)); doc_table.style = "Light Shading Accent 1"
            headers = ([table.index.name or ""] if doc_table.columns.__len__() == len(table.columns)+1 else []) + [str(c) for c in table.columns]
            for cell, text in zip(doc_table.rows[0].cells, headers): cell.text = text
            for index, row in table.iterrows():
                values = ([str(index)] if len(headers) == len(table.columns)+1 else []) + [str(v) for v in row.tolist()]
                cells = doc_table.add_row().cells
                for cell, text in zip(cells, values): cell.text = text
            document.add_paragraph()
        for fig in result.get("figures", []):
            image = BytesIO(); fig.savefig(image, format="png", dpi=180, bbox_inches="tight", facecolor="white"); image.seek(0); document.add_picture(image, width=Inches(6.2))
        if result.get("interpretation"): document.add_paragraph(result["interpretation"])
    output = BytesIO(); document.save(output); return output.getvalue()
